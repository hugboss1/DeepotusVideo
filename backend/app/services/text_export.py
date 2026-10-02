"""Tâche #65 (plan chapitres T15-T16, 02/10/2026) — les EXPORTS d'un chapitre : manuscrit (.docx, .pdf), scénario
(.docx, .pdf), storyboard (.pdf). Rendus en OCTETS (aucun fichier temporaire), aucun réseau, aucune dépense.

DÉCISIONS DE L'UTILISATEUR (02/10) : A4, sauf le scénario au format US Letter (le standard du métier) ; PDF composés
sur pypdf (pdf_compose) ; ce qui sort de la police standard est remplacé et COMPTÉ (le .docx garde tout) ; le
storyboard en quatre cartes verticales par page, la vignette 9:16 ENTIÈRE (production, sinon croquis, étiquetée).
Le scénario est classé par `screenplay_import.elements` — la grammaire de l'import, pas une seconde.
"""
from __future__ import annotations

import io
from pathlib import Path

from app.services import pdf_compose as PC
from app.services import screenplay_import as SI


def _paragraphes(texte: str) -> list[list[str]]:
    """Paragraphes (séparés par une ligne vide), chacun en LIGNES (un saut simple est gardé : « — Bonjour. »
    puis « — Salut. » restent deux lignes, comme dans le .docx)."""
    out, cur = [], []
    for l in (texte or "").replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        if l.strip():
            cur.append(l.rstrip())
        elif cur:
            out.append(cur)
            cur = []
    if cur:
        out.append(cur)
    return out


def duree_fr(s: float | None) -> str:
    v = float(s or 0)
    if v < 60:
        return f"{v:.1f}".replace(".", ",") + " s"
    m, r = divmod(int(v + 0.5), 60)          # arrondi au plus proche (round() arrondit 86,5 à 86)
    return f"{m} min {r:02d} s"


# ─────────────────────────────────────────────── manuscrit ───────────────────────────────────────────────

def manuscrit_docx(titre: str, texte: str) -> bytes:
    import docx
    from docx.shared import Cm, Pt
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for cote in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, cote, Cm(2.5))
    d.styles["Normal"].font.size = Pt(11.5)
    d.add_heading(titre or "Chapitre", level=0)
    for par in _paragraphes(texte):
        p = d.add_paragraph("\n".join(par))
        p.paragraph_format.space_after = Pt(8)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def manuscrit_pdf(titre: str, texte: str) -> tuple[bytes, int]:
    doc = PC.Document(PC.A4)
    marge, haut, bas, corps, interligne = 72.0, 72.0, PC.A4[1] - 72.0, 11.0, 15.5
    utile = PC.A4[0] - 2 * marge
    y = haut
    for l in PC.couper(titre or "Chapitre", "HB", 18, utile):
        doc.texte(marge, y + 18, l, "HB", 18)
        y += 24
    y += 14

    def saut():
        nonlocal y
        doc.nouvelle_page()
        y = haut

    for par in _paragraphes(texte):
        for ligne in par:
            for morceau in PC.couper(ligne, "H", corps, utile):
                if y + interligne > bas:
                    saut()
                doc.texte(marge, y + corps, morceau, "H", corps)
                y += interligne
        y += interligne * 0.6
    for i in range(len(doc.pages)):
        doc.pages[i]["ops"].append(f"BT 0.45 g /H 8.5 Tf {PC._n(PC.A4[0] / 2 - 4)} 36 Td ({i + 1}) Tj ET 0 g".encode("ascii"))
    return doc.octets(titre or "Chapitre"), doc.remplacements


# ─────────────────────────────────────────────── scénario ───────────────────────────────────────────────
# US Letter, Courier 12 : marge gauche 1,5 po, droite 1 po. Retraits mesurés DEPUIS LE BORD de la page (PDF) et
# depuis la marge (docx) — les deux exports tombent au même endroit.
_PO = 72.0
_RETRAITS = {  # type -> (gauche depuis le bord, largeur), en pouces
    "action": (1.5, 6.0), "en_tete": (1.5, 6.0), "personnage": (3.7, 3.3), "dialogue": (2.5, 3.5),
    "parenthese": (3.1, 2.3), "centre": (1.5, 6.0), "transition": (1.5, 6.0),
}


def _scenario_elements(scenes: list[dict]) -> list[tuple[str, str]]:
    out: list[tuple[str, str]] = []
    for s in scenes:
        out.append(("en_tete", (s.get("slugline") or "").upper()))
        out.append(("vide", ""))
        out += SI.elements(s.get("fountain_text") or "")
        out.append(("vide", ""))
    return out


def scenario_pdf(titre: str, scenes: list[dict]) -> tuple[bytes, int]:
    L, H = PC.LETTER
    doc = PC.Document(PC.LETTER)
    # page de titre
    tl = PC.couper((titre or "Scénario").upper(), "C", 12, 6 * _PO)
    y0 = H / 2 - 12 * len(tl)
    for k, l in enumerate(tl):
        doc.texte((L - PC.largeur(l, "C", 12)) / 2, y0 + 12 * k, l, "C", 12)
    doc.nouvelle_page()
    haut, bas, il = _PO, H - _PO, 12.0
    y = haut
    pages = 1

    def saut():
        nonlocal y, pages
        doc.nouvelle_page()
        pages += 1
        doc.texte_droite(L - _PO, _PO / 2 + 12, f"{pages}.", "C", 12)
        y = haut

    for ty, tx in _scenario_elements(scenes):
        if ty == "vide":
            y += il
            continue
        gauche, larg = _RETRAITS[ty]
        police = "CB" if ty == "en_tete" else "C"
        for ligne in PC.couper(tx, police, 12, larg * _PO):
            if y + il > bas:
                saut()
            if ty == "transition":
                doc.texte_droite(L - _PO, y + 10, ligne.upper(), "C", 12)
            elif ty == "centre":
                doc.texte((L - PC.largeur(ligne, "C", 12)) / 2, y + 10, ligne, "C", 12)
            else:
                doc.texte(gauche * _PO, y + 10, ligne.upper() if ty in ("personnage", "en_tete") else ligne, police, 12)
            y += il
    return doc.octets(titre or "Scénario"), doc.remplacements


def scenario_docx(titre: str, scenes: list[dict]) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Inches(1.5), Inches(1), Inches(1), Inches(1)
    st = d.styles["Normal"]
    st.font.name, st.font.size = "Courier New", Pt(12)
    st.paragraph_format.space_after, st.paragraph_format.space_before = Pt(0), Pt(0)
    t = d.add_paragraph((titre or "Scénario").upper())
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    d.add_page_break()
    utile = 6.0
    for ty, tx in _scenario_elements(scenes):
        p = d.add_paragraph("" if ty == "vide" else (tx.upper() if ty in ("personnage", "en_tete", "transition") else tx))
        if ty == "vide":
            continue
        gauche, larg = _RETRAITS[ty]
        pf = p.paragraph_format
        pf.left_indent = Inches(gauche - 1.5)
        pf.right_indent = Inches(max(0.0, utile - (gauche - 1.5) - larg))
        if ty == "en_tete" and p.runs:
            p.runs[0].bold = True
        if ty == "transition":
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if ty == "centre":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ─────────────────────────────────────────────── storyboard ───────────────────────────────────────────────

def storyboard_pdf(titre: str, shots: list[dict], images: Path, noms: dict[str, str] | None = None) -> tuple[bytes, int]:
    """Quatre cartes par page A4 (2×2). Vignette 9:16 ENTIÈRE à gauche (production, sinon croquis — étiquetée),
    la fiche du plan à droite, l'action en pleine largeur dessous."""
    noms = noms or {}
    L, H = PC.A4
    doc = PC.Document(PC.A4)
    marge, entete, gout = 36.0, 46.0, 12.0
    cw = (L - 2 * marge - gout) / 2
    ch = (H - 2 * marge - entete - gout) / 2
    total = sum(float(s.get("duration_s") or 0) for s in shots)
    racine = Path(images).resolve()

    def tete(n_page: int):
        for k, l in enumerate(PC.couper(titre or "Storyboard", "HB", 14, L - 2 * marge - 170)[:2]):
            doc.texte(marge, marge + 14 + 16 * k, l, "HB", 14)
        doc.texte_droite(L - marge, marge + 12, f"{len(shots)} plans · {duree_fr(total)} · page {n_page}", "H", 8.5)
        doc.trait(marge, marge + entete - 10, L - marge, marge + entete - 10, 0.6, 0.7)

    tete(1)
    for k, s in enumerate(shots):
        if k and k % 4 == 0:
            doc.nouvelle_page()
            tete(k // 4 + 1)
        col, rang = (k % 4) % 2, (k % 4) // 2
        x = marge + col * (cw + gout)
        y = marge + entete + rang * (ch + gout)
        doc.cadre(x, y, cw, ch, 0.8)
        iw = 112.0
        ih = iw * 16 / 9
        nom = s.get("image") or s.get("sketch_image")
        sorte = "production" if s.get("image") else "croquis" if s.get("sketch_image") else None
        posee = None
        if nom:
            chemin = racine / Path(str(nom)).name          # .name : un « ../ » ne sort jamais du dossier des images
            if chemin.is_file():
                posee = doc.image(chemin, x + 8, y + 8, iw, ih)
        doc.cadre(x + 8, y + 8, iw, ih, 0.85)
        if posee is None:
            sorte = "aucune image"
            doc.texte(x + 8 + (iw - PC.largeur(sorte, "HI", 8)) / 2, y + 8 + ih / 2, sorte, "HI", 8, gris=0.5)
        else:
            doc.texte(x + 8, y + 8 + ih + 10, sorte, "HI", 7, gris=0.45)
        tx, tw = x + 8 + iw + 10, cw - iw - 26
        yy = y + 20
        doc.texte(tx, yy, f"PLAN {int(s.get('idx', k)) + 1}", "HB", 11)
        yy += 16
        fiche = [("Cadrage", s.get("shot_type")), ("Caméra", s.get("camera_move")), ("Durée", duree_fr(s.get("duration_s"))),
                 ("Mouvement", s.get("motion_recipe"))]
        for lib, val in fiche:
            if not val:
                continue
            doc.texte(tx, yy, lib, "HB", 7.5, gris=0.4)
            yy += 10
            for l in PC.couper(str(val), "H", 8.5, tw)[:2]:
                doc.texte(tx, yy, l, "H", 8.5)
                yy += 11
            yy += 3
        ents = [noms.get(e, "") for e in (s.get("entities") or []) if noms.get(e)]
        if ents:
            doc.texte(tx, yy, "Entités", "HB", 7.5, gris=0.4)
            yy += 10
            for l in PC.couper(", ".join(ents), "H", 8.5, tw)[:4]:
                doc.texte(tx, yy, l, "H", 8.5)
                yy += 11
        # l'action, en pleine largeur sous la vignette ; coupée avec « … » si elle déborde de la carte
        action = (s.get("action") or s.get("source_text") or "").strip()
        ay, fin = y + 8 + ih + 24, y + ch - 8
        lignes = PC.couper(action, "H", 9, cw - 16) if action else []
        place = max(0, int((fin - ay) // 11.5))
        if len(lignes) > place:
            lignes = lignes[:place]
            if lignes:
                lignes[-1] = lignes[-1].rstrip(" .,;") + " …"
        for l in lignes:
            doc.texte(x + 8, ay + 9, l, "H", 9)
            ay += 11.5
    return doc.octets(f"{titre or 'Chapitre'} — storyboard"), doc.remplacements
