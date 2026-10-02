# -*- coding: utf-8 -*-
"""Plan chapitres T14-T16 (tache #65 du suivi, 02/10/2026) — les EXPORTS d'un chapitre : manuscrit .docx/.pdf,
scenario .docx/.pdf, storyboard .pdf. Aucune depense. Chaque PDF est RELU par pypdf (strict), chaque .docx par
python-docx : on verifie ce qui est ecrit, pas ce qu'on croit avoir ecrit.
Ce que le plan faisait faux, et que ce banc garde :
  - vignettes 9:16 ecrasees ×3 dans une case paysage -> l'image ENTIERE, proportions gardees ;
  - le classement du scenario (# ! @ > (cont'd) [[ ]]) -> la grammaire de l'import, une seule ;
  - « ł » et les emoji devenaient « ? » sans le dire -> remplaces au plus proche ET comptes (X-DZ-Remplacements) ;
  - le titre des metadonnees en mojibake (cp1252 brut) -> encode par pypdf ;
  - la largeur des caracteres au-dela d'ASCII (— = 1000/1000) ; l'alpha noirci ; les images en double ;
  - les erreurs 400 telechargees comme un fichier ; le .fountain qui plantait (500) sur un titre avec « ’ ».
DECISIONS DE L'UTILISATEUR (02/10) : pypdf comme Card Forge ; police standard + repli dit ; 4 cartes 9:16 par page A4 ;
les cinq exports, A4 sauf le scenario en US Letter.
Temoin positif : la base (3625f530) n'a ni les modules ni les routes.
Run (depuis backend/) : & $PY tests/test_chapitres_exports.py"""
import io, json, os, pathlib, re, sqlite3, subprocess, sys, tempfile, types
sys.stdout.reconfigure(encoding="utf-8", errors="replace")
_tmp = pathlib.Path(tempfile.mkdtemp(prefix="dzexport_"))
os.environ["DEEPOTUS_DATA_DIR"] = str(_tmp)
_DB = _tmp / "t.db"
os.environ["DATABASE_URL"] = f"sqlite+aiosqlite:///{_DB.as_posix()}"
for d in ("images", "outputs"):
    (_tmp / d).mkdir(exist_ok=True)
os.environ["IMAGES_FOLDER"] = str(_tmp / "images")
os.environ["OUTPUTS_FOLDER"] = str(_tmp / "outputs")
for k in ("ANTHROPIC_API_KEY", "OPENAI_API_KEY", "OLLAMA_URL", "GEMINI_API_KEY", "OPENROUTER_API_KEY", "FAL_KEY",
          "ELEVENLABS_API_KEY"):
    os.environ[k] = ""
_ICI = pathlib.Path(__file__).resolve().parent
sys.path.insert(0, str(_ICI.parent))
from loguru import logger                                           # noqa: E402
logger.remove()

ok = fail = 0
def check(label, cond, detail=""):
    global ok, fail
    if cond: ok += 1; print(f"  PASS  {label}")
    else: fail += 1; print(f"  FAIL  {label} {detail}")


BASE = "3625f530"
r0 = subprocess.run(["git", "show", f"{BASE}:backend/app/services/text_export.py"], capture_output=True, cwd=str(_ICI.parent))
r1 = subprocess.run(["git", "show", f"{BASE}:backend/app/api/routes.py"], capture_output=True, cwd=str(_ICI.parent))
check("T1 temoin : la base n'a ni les exports ni leurs routes", r0.returncode != 0 and r1.returncode == 0
      and b'"/chapters/{chapter_id}/storyboard.pdf"' not in r1.stdout and b'"/chapters/{chapter_id}/export.pdf"' not in r1.stdout)

from PIL import Image                                               # noqa: E402
from pypdf import PdfReader                                         # noqa: E402
try:
    from app.services import pdf_compose as PC                      # noqa: E402
    from app.services import text_export as TE                      # noqa: E402
except ImportError as e:
    PC = TE = None
    check("T2 les modules existent", False, str(e))


def lire(b):
    try:
        return PdfReader(io.BytesIO(b), strict=True)
    except Exception as e:
        print("   (pdf illisible)", e)
        return None


def positions(page):
    """[(texte, x, y)] des fragments poses, lus par pypdf."""
    out = []
    page.extract_text(visitor_text=lambda t, cm, tm, fd, fs: out.append((t, tm[4] * cm[0] + cm[4], tm[5] * cm[3] + cm[5])) if t.strip() else None)
    return out


def images_posees(page):
    """[(largeur, hauteur)] des images posees (operateur cm suivi de Do)."""
    data = page.get_contents().get_data().decode("latin-1")
    return [(float(a), float(d)) for a, d in re.findall(r"q ([\d.]+) 0 0 ([\d.]+) [\d.]+ [\d.]+ cm /Im\w+ Do Q", data)]


if PC:
    print("[C] le compositeur")
    b, n = PC.winansi("L’Éveil « x » — … œ € ł 🎬")
    check("C1 ’ « » — … œ € passent ; ł -> l, l'emoji -> ? : DEUX remplacements comptes", n == 2 and b.decode("cp1252") == "L’Éveil « x » — … œ € l ?", repr(b))
    check("C2 largeurs mesurees : n = 556, — = 1000 (Helvetica) ; Courier 600 pour tout",
          PC.largeur("n", "H", 1000) == 556 and PC.largeur("—", "H", 1000) == 1000 and PC.largeur("—ł", "C", 1000) == 1200)
    lignes = PC.couper("un — deux — trois " * 20 + "anticonstitutionnellementanticonstitutionnellement", "H", 11, 120)
    check("C3 couper : aucune ligne plus large que la place, meme un mot trop long", all(PC.largeur(l, "H", 11) <= 120 for l in lignes) and len(lignes) > 5)
    t = _tmp / "img"
    t.mkdir()
    Image.new("RGBA", (90, 160), (0, 0, 0, 0)).save(t / "transp.png")
    Image.new("RGB", (540, 960), (200, 40, 40)).save(t / "prod.png")
    d = PC.Document()
    d.texte(72, 100, "L’Éveil « Wałkuski » — fin")
    a = d.image(t / "transp.png", 72, 120, 90, 160)
    d.image(t / "transp.png", 200, 120, 90, 160)
    boite = d.image(t / "prod.png", 300, 120, 112, 112)
    d.nouvelle_page()
    d.texte(72, 100, "page deux", "C", 12)
    r = lire(d.octets("L’Éveil — storyboard"))
    check("C4 relu en STRICT : deux pages, le titre des metadonnees EXACT (’ et —), le texte extrait",
          r is not None and len(r.pages) == 2 and r.metadata.title == "L’Éveil — storyboard"
          and "L’Éveil « Walkuski » — fin" in r.pages[0].extract_text(), repr(r.metadata.title) if r else "")
    xo = r.pages[0]["/Resources"]["/XObject"] if r else {}
    imgs = [xo[k].get_object() for k in xo] if r else []
    blanc = None
    for o in imgs:
        if o["/Width"] < 200:
            blanc = Image.open(io.BytesIO(o._data)).convert("RGB").getpixel((45, 80))
    check("C5 une image identique n'est ecrite qu'UNE fois ; l'alpha est compose sur BLANC (pas noirci)",
          len(imgs) == 2 and blanc is not None and min(blanc) > 240, f"{len(imgs)} {blanc}")
    check("C6 l'image ENTIERE dans sa boite : 540×960 dans 112×112 -> 63×112, centree", boite is not None
          and abs(boite[2] - 63.0) < 0.1 and abs(boite[3] - 112) < 0.1 and abs(boite[0] - (300 + (112 - 63) / 2)) < 0.1, str(boite))
    check("C7 les images sont en JPEG (DCTDecode)", all(o["/Filter"] == "/DCTDecode" for o in imgs))

    print("\n[S] le scenario")
    from app.services import screenplay_import as SI
    el = SI.elements("# BOARD\n\n!BOOM.\n\nUNE ACTION EN CAPITALES.\n\nVANE (cont'd)\n(bas)\nIl pleut.\n\n@McGregor\nOui.\n\n[[note]]\n> FADE OUT.\n\n>FIN<\n\nCUT TO:")
    check("S1 une seule grammaire : section ecartee, « ! » action, capitales seules = action, (cont'd), @, notes, >, ><, CUT TO:",
          [e for e in el if e[0] != "vide"] == [("action", "BOOM."), ("action", "UNE ACTION EN CAPITALES."), ("personnage", "VANE (cont'd)"),
          ("parenthese", "(bas)"), ("dialogue", "Il pleut."), ("personnage", "MCGREGOR"), ("dialogue", "Oui."),
          ("transition", "FADE OUT."), ("centre", "FIN"), ("transition", "CUT TO:")], str(el))
    sc = [{"slugline": "INT. CUISINE - NUIT", "fountain_text": "Vane entre.\n\nVANE\n(bas)\nIl pleut.\n\nCUT TO:"}] + \
         [{"slugline": f"EXT. RUE {i} - JOUR", "fountain_text": "Une action. " * 30} for i in range(12)]
    b, n = TE.scenario_pdf("Le Phare", sc)
    r = lire(b)
    pos = {t.strip(): (round(x, 1), round(y, 1)) for t, x, y in positions(r.pages[1])} if r else {}
    check("S2 US Letter, page de titre, puis les retraits du metier : action 1,5 po, replique 3,7 po, dialogue 2,5 po, transition a droite",
          r is not None and [float(v) for v in r.pages[0].mediabox[2:]] == [612, 792] and "LE PHARE" in r.pages[0].extract_text()
          and pos.get("Vane entre.", (0,))[0] == 108 and pos.get("VANE", (0,))[0] == 266.4 and pos.get("Il pleut.", (0,))[0] == 180
          and pos.get("(bas)", (0,))[0] == 223.2 and pos.get("CUT TO:", (0,))[0] == 540 - 7 * 7.2, str(pos))
    check("S3 pages numerotees a la maniere du metier : ni la page de titre ni la 1re page du texte, « 2. » ensuite",
          r is not None and len(r.pages) >= 3 and not re.search(r"(^|\n)1\.", r.pages[1].extract_text())
          and re.search(r"(^|\n)2\.", r.pages[2].extract_text()) is not None, r.pages[2].extract_text()[:80] if r else "")
    import docx as _docx
    dd = _docx.Document(io.BytesIO(TE.scenario_docx("Le Phare", sc[:1])))
    pars = {p.text: p for p in dd.paragraphs if p.text}
    sec = dd.sections[0]
    check("S4 .docx : Letter, marge gauche 1,5 po, Courier New ; replique a +2,2 po, dialogue +1,0 po (droite 1,5 po) — comme le PDF",
          round(sec.page_width.inches, 2) == 8.5 and round(sec.left_margin.inches, 2) == 1.5 and dd.styles["Normal"].font.name == "Courier New"
          and round(pars["VANE"].paragraph_format.left_indent.inches, 2) == 2.2 and round(pars["Il pleut."].paragraph_format.left_indent.inches, 2) == 1.0
          and round(pars["Il pleut."].paragraph_format.right_indent.inches, 2) == 1.5, str(list(pars)))

    print("\n[M] le manuscrit")
    b, n = TE.manuscrit_pdf("Titre", "— Bonjour.\n— Salut.\n\nUn paragraphe.\n\n" + ("Long texte. " * 400))
    r = lire(b)
    check("M1 A4, les sauts SIMPLES gardes (deux repliques = deux lignes), pages numerotees",
          r is not None and [round(float(v)) for v in r.pages[0].mediabox[2:]] == [595, 842] and "— Bonjour.\n— Salut." in r.pages[0].extract_text()
          and len(r.pages) >= 2 and r.pages[1].extract_text().rstrip().endswith("2"), r.pages[0].extract_text()[:60] if r else "")
    xs = [x for p in (r.pages if r else []) for t, x, y in positions(p) if t.strip() and len(t) > 20]
    check("M2 marges tenues : aucune ligne ne commence hors de la marge de 72 pt", xs and all(abs(x - 72) < 0.5 for x in xs), str(set(round(x) for x in xs)))
    dd = _docx.Document(io.BytesIO(TE.manuscrit_docx("Titre", "— Bonjour.\n— Salut.\n\nDeux.")))
    check("M3 .docx : A4, titre, et les sauts simples gardes", round(dd.sections[0].page_width.cm, 1) == 21.0
          and dd.paragraphs[0].text == "Titre" and any(p.text == "— Bonjour.\n— Salut." for p in dd.paragraphs), str([p.text for p in dd.paragraphs]))

    print("\n[B] le storyboard")
    Image.new("RGB", (540, 960), (40, 120, 200)).save(t / "croquis.png")
    (_tmp / "secret.png").write_bytes((t / "prod.png").read_bytes())
    shots = [{"idx": 0, "shot_type": "Gros plan", "camera_move": "Travelling", "duration_s": 4.0, "action": "Vane entre.", "entities": ["e1"],
              "image": "prod.png", "sketch_image": "croquis.png"},
             {"idx": 1, "shot_type": "Plan large", "duration_s": 3.5, "action": "Une longue action. " * 60, "sketch_image": "croquis.png"},
             {"idx": 2, "duration_s": 2, "source_text": "Sans image."},
             {"idx": 3, "duration_s": 2, "action": "Evasion.", "image": "../secret.png"},
             {"idx": 4, "duration_s": 75, "action": "Cinquieme."}]
    b, n = TE.storyboard_pdf("L’Éveil", shots, t, {"e1": "Vane Ardel"})
    r = lire(b)
    p0 = r.pages[0].extract_text() if r else ""
    poses = images_posees(r.pages[0]) if r else []
    check("B1 la vignette 9:16 est ENTIERE (ratio 0,5625 garde), pas ecrasee dans une case paysage",
          len(poses) == 2 and all(abs(w / h - 0.5625) < 0.01 for w, h in poses) and poses[0][1] > 190, str(poses))
    check("B2 quatre cartes par page A4 : cinq plans -> deux pages ; le titre avec ’ et — dans les metadonnees",
          r is not None and len(r.pages) == 2 and r.metadata.title == "L’Éveil — storyboard" and "PLAN 5" in r.pages[1].extract_text())
    check("B3 production d'abord, sinon croquis, sinon « aucune image » — etiquetes", p0.index("production") < p0.index("PLAN 1") < p0.index("croquis")
          and p0.count("aucune image") == 2, p0[:200])
    check("B4 une image HORS du dossier des images (../) n'est jamais lue", "Evasion." in p0 and len(poses) == 2)
    ordre = re.findall(r"cm /(Im\w+) Do", r.pages[0].get_contents().get_data().decode("latin-1")) if r else []
    xo0 = r.pages[0]["/Resources"]["/XObject"] if r else {}
    couleurs = [Image.open(io.BytesIO(xo0["/" + k].get_object()._data)).convert("RGB").getpixel((10, 10)) for k in ordre]
    check("B7 le plan qui a LES DEUX images montre la PRODUCTION (rouge), pas le croquis (bleu) ; le plan 2, son croquis",
          len(couleurs) == 2 and couleurs[0][0] > 150 and couleurs[0][2] < 90 and couleurs[1][2] > 150, str(couleurs))
    check("B5 la fiche : cadrage, camera, duree a la francaise, entites par leur NOM ; l'action trop longue coupee avec « … »",
          "Gros plan" in p0 and "Travelling" in p0 and "4,0 s" in p0 and "Vane Ardel" in p0 and "…" in p0 and "1 min 15 s" in r.pages[1].extract_text()
          and "86,5 s" not in p0, p0[:300])
    check("B6 l'en-tete : nombre de plans et duree totale", "5 plans · 1 min 27 s · page 1" in p0, p0[:120])

# ── les routes ────────────────────────────────────────────────────────────────────────────────────────────────────
sys.modules["fal_client"] = types.ModuleType("fal_client")
from fastapi.testclient import TestClient                           # noqa: E402
from app.main import app                                            # noqa: E402
import app.main as _MAIN                                            # noqa: E402


async def _boucle_coupee():
    return None
_MAIN.schedule_loop = _boucle_coupee


def js(r):
    try:
        v = r.json()
        return v if isinstance(v, dict) else {"_liste": v}
    except Exception:
        return {}


print("\n[R] les routes")
with TestClient(app, client=("127.0.0.1", 50000), raise_server_exceptions=False) as c:
    check("R1 chapitre inconnu : 404 partout", all(c.get(f"/api/chapters/inconnu/{p}").status_code == 404
          for p in ("export.docx", "export.pdf?kind=scenario", "storyboard.pdf")))
    vide = js(c.post("/api/chapters", json={"title": "Vide", "script_text": ""}))
    vid = vide.get("id")
    r = c.get(f"/api/chapters/{vid}/export.pdf")
    check("R2 manuscrit vide : 400 qui le DIT", r.status_code == 400 and "vide" in r.text, r.text[:120])
    check("R3 kind inconnu : 400", c.get(f"/api/chapters/{vid}/export.pdf?kind=roman").status_code == 400)
    r = c.get(f"/api/chapters/{vid}/export.docx?kind=scenario")
    check("R4 sans scenario : 400 qui dit quoi faire (Adapter / Importer)", r.status_code == 400 and "Adapter" in r.text and "Importer" in r.text)
    r = c.get(f"/api/chapters/{vid}/storyboard.pdf")
    check("R5 sans storyboard : 400 qui dit quoi faire", r.status_code == 400 and "découpe" in r.text)
    ch = js(c.post("/api/chapters", json={"title": "L’Éveil — Wałkuski", "script_text": "Vane entre sous la pluie.\n\nYsolde ł se tait."}))
    cid = ch.get("id")
    r = c.get(f"/api/chapters/{cid}/export.docx")
    check("R6 .docx du manuscrit : 200, le bon type, un nom qui survit a « ’ » (filename* UTF-8) et un repli ASCII",
          r.status_code == 200 and r.headers.get("content-type", "").startswith("application/vnd.openxmlformats")
          and "filename*=UTF-8''L%E2%80%99%C3%89veil%20%E2%80%94%20Wa%C5%82kuski.docx" in r.headers.get("content-disposition", "")
          and 'filename="LEveil Walkuski.docx"' in r.headers.get("content-disposition", ""), r.headers.get("content-disposition"))
    r = c.get(f"/api/chapters/{cid}/export.pdf")
    rr = lire(r.content) if r.status_code == 200 else None
    check("R7 .pdf du manuscrit : relu ; les caracteres remplaces (ł ×2 : titre + texte) COMPTES dans X-DZ-Remplacements",
          rr is not None and r.headers.get("x-dz-remplacements") == "2" and "Ysolde l se tait." in rr.pages[0].extract_text(),
          f"{r.status_code} {r.headers.get('x-dz-remplacements')}")
    r = c.get(f"/api/chapters/{cid}/screenplay", params={"format": "fountain"})
    check("R8 le .fountain d'un titre avec « ’ » ne PLANTE plus (500 avant)", r.status_code == 200
          and "filename*=UTF-8''" in r.headers.get("content-disposition", ""), f"{r.status_code}")
    c.post(f"/api/chapters/{cid}/screenplay/import", files={"file": ("s.fountain", b"INT. CUISINE - NIGHT\n\nVane entre.\n\nVANE\nIl pleut.")})
    r = c.get(f"/api/chapters/{cid}/export.pdf?kind=scenario")
    rr = lire(r.content) if r.status_code == 200 else None
    check("R9 .pdf du scenario : US Letter, la scene et sa replique", rr is not None and [float(v) for v in rr.pages[0].mediabox[2:]] == [612, 792]
          and "INT. CUISINE - NUIT" in rr.pages[1].extract_text() and "Il pleut." in rr.pages[1].extract_text(), f"{r.status_code}")
    shots = js(c.post(f"/api/chapters/{cid}/storyboard/decoupe", json={"method": "paragraph"})).get("shots", [])
    Image.new("RGB", (540, 960), (90, 90, 160)).save(_tmp / "images" / "croquis_banc.png")
    with sqlite3.connect(str(_DB)) as db:
        db.execute("UPDATE shots SET sketch_image='croquis_banc.png' WHERE id=?", (shots[0]["id"],))
    r = c.get(f"/api/chapters/{cid}/storyboard.pdf")
    rr = lire(r.content) if r.status_code == 200 else None
    check("R10 le storyboard : 200, deux plans, la vignette du croquis posee, nom « … - storyboard.pdf »",
          rr is not None and "PLAN 2" in rr.pages[0].extract_text() and len(images_posees(rr.pages[0])) == 1
          and "storyboard.pdf" in r.headers.get("content-disposition", ""), f"{r.status_code}")

print("\n[U] /atelier (cablage ; comportement prouve dans le navigateur)")
_at = _ICI.parent.parent / "frontend" / "atelier"
_js = (_at / "atelier.js").read_text(encoding="utf-8")
_html = (_at / "index.html").read_text(encoding="utf-8")
_fn = _js.split("async function telechargerExport(")[1].split("\n}\n")[0] if "async function telechargerExport(" in _js else ""
boutons = re.findall(r'<button id="exp\w+" class="btn" data-export="([^"]+)" title="[^"]+">', _html)
check("U1 cinq boutons d'export, chacun avec un title", sorted(boutons) == sorted(["export.docx?kind=manuscrit", "export.pdf?kind=manuscrit",
      "export.docx?kind=scenario", "export.pdf?kind=scenario", "storyboard.pdf"]), str(boutons))
check("U2 telecharges par fetch : une erreur passe par le dialogue maison (pas un fichier JSON), les remplacements sont DITS",
      "if (!r.ok)" in _fn and "window.__dzDialogue.informer(" in _fn and "X-DZ-Remplacements" in _fn and "filename\\*=UTF-8''" in _fn
      and '"[data-export]"' in _js and "telechargerExport(b.dataset.export)" in _js)

print(f"\n{ok} PASS / {fail} FAIL")
sys.exit(1 if fail else 0)
